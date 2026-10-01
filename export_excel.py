"""Create a portable Excel summary from the processed JSON results."""
from pathlib import Path
import argparse, json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEAD_FILL = PatternFill("solid", fgColor="254E70")
HEAD_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10, color="22364B")


def format_sheet(ws):
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 36
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column in range(1, ws.max_column + 1):
        values = [str(ws.cell(row, column).value or "") for row in range(1, min(ws.max_row, 60) + 1)]
        ws.column_dimensions[get_column_letter(column)].width = min(48, max(11, max(map(len, values)) + 2))
    ws.auto_filter.ref = ws.dimensions


def main():
    parser = argparse.ArgumentParser(description="Xuất Excel tổng hợp PETG")
    parser.add_argument("--output", type=Path, required=True, help="Thư mục kết quả đã xử lý")
    args = parser.parse_args()
    rows = json.loads((args.output / "analysis.json").read_text(encoding="utf-8"))["results"]
    groups = json.loads((args.output / "groups.json").read_text(encoding="utf-8"))

    wb = Workbook()
    summary = wb.active
    summary.title = "84 trường hợp"
    summary.append(["Trường hợp", "Góc in (°)", "Lớp (mm)", "Infill (%)", "Pattern", "n", "UTS TB (MPa)", "SD (MPa)", "Min", "Max", "CV (%)", "n đường qua QA", "Lưu ý"])
    for g in groups:
        summary.append([g["group"], g["angle_deg"], g["layer_mm"], g["infill_pct"], g["pattern"], g["n"], g["UTS_raw_mean_MPa"], g["UTS_raw_sd_MPa"], g["UTS_raw_min_MPa"], g["UTS_raw_max_MPa"], g["UTS_raw_cv_pct"], g["n_strain_QA"], g["group_status"]])

    samples = wb.create_sheet("420 mẫu")
    samples.append(["Mẫu", "Trường hợp", "Lặp", "Góc in (°)", "Lớp (mm)", "Infill (%)", "Pattern", "Fmax (kgf)", "Fmax (N)", "UTS (MPa)", "Nguồn chuyển vị", "ε tại UTS (%)", "Số điểm", "Trạng thái", "Cờ kiểm tra", "Ghi chú"])
    for r in rows:
        samples.append([r["sample"], r["group"], r["replicate"], r["angle_deg"], r["layer_mm"], r["infill_pct"], r["pattern"], r["Fmax_terminal_raw_kgf"], r["Fmax_terminal_raw_N"], r["UTS_raw_MPa"], r.get("displacement_source"), r.get("eps_UTS_candidate_pct"), r["candidate_curve_points"], r["status"], "; ".join(r["quality_flags"]), "; ".join(r["notes"])])

    qa = wb.create_sheet("Đồng bộ và QA")
    qa.append(["Mẫu", "Phương pháp", "Nguồn thời gian", "Scale", "Offset (s)", "RMSE", "Lệch mốc max (s)", "Cắt (s)", "Xuống 50 kg", "Nguồn chuyển vị", "Cờ kiểm tra"])
    for r in rows:
        qa.append([r["sample"], r.get("sync_method"), r.get("time_basis"), r.get("time_scale"), r.get("offset_s"), r.get("rmse_normalized"), r.get("max_landmark_residual_s"), r.get("last_valid_time_s"), r.get("terminal_below50_confirmed"), r.get("displacement_source"), "; ".join(r["quality_flags"])])

    for sheet in wb.worksheets:
        format_sheet(sheet)
    target = args.output / "Tong_hop_420_mau.xlsx"
    wb.save(target)
    print(target)


if __name__ == "__main__":
    main()
